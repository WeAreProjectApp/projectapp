"""Read-only, local extraction of bounded PDF, DOCX and XLSX attachments."""
from contextlib import contextmanager
from io import BytesIO
import json
from pathlib import Path
import subprocess
import sys
from zipfile import ZipFile

from defusedxml import ElementTree

from content.services.markdown_export import (
    MarkdownBuffer, MarkdownExportError, export_payload, literal, table,
)

MAX_FILE_BYTES = 15 * 1024 * 1024
MAX_EXPANDED_BYTES = 50 * 1024 * 1024
MAX_PDF_PAGES = 100
MAX_SOURCE_CHARACTERS = 60_000
MAX_SOURCE_FRAGMENTS = 1000
MAX_SOURCE_SECTION_CHARACTERS = 200
MAX_SHEET_CELLS = 20_000
WORD_NS = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
PDF_WARNING = 'Texto extraído del PDF; el formato fue reconstruido. Las imágenes y firmas gráficas no se copian.'
OFFICE_WARNING = 'Vista del contenido textual; el diseño, las imágenes y los objetos incrustados no se reproducen.'


def read_attachment(document):
    if not document.file:
        raise MarkdownExportError('El archivo ya no está disponible.', 'file_missing', 404)
    try:
        with document.file.open('rb') as source:
            data = source.read(MAX_FILE_BYTES + 1)
    except (OSError, ValueError) as exc:
        raise MarkdownExportError('El archivo ya no está disponible.', 'file_missing', 404) from exc
    if len(data) > MAX_FILE_BYTES:
        raise MarkdownExportError('El archivo supera el límite de 15 MB.', 'document_too_large', 413)
    return data


@contextmanager
def office_archive(data):
    with ZipFile(BytesIO(data)) as archive:
        entries = archive.infolist()
        if len(entries) > 2000 or sum(entry.file_size for entry in entries) > MAX_EXPANDED_BYTES:
            raise MarkdownExportError('El archivo supera el límite de contenido descomprimido.', 'document_too_large', 413)
        if any(entry.flag_bits & 1 for entry in entries):
            raise MarkdownExportError('El archivo está protegido. Descarga una copia sin contraseña.', 'document_protected')
        yield archive


def pdf_markdown(data):
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise MarkdownExportError('El PDF está protegido. Usa una copia sin contraseña.', 'document_protected')
    if len(reader.pages) > MAX_PDF_PAGES:
        raise MarkdownExportError('El PDF supera el límite de 100 páginas.', 'document_too_large', 413)
    output = MarkdownBuffer()
    blank_pages = []
    expanded = 0
    for number, page in enumerate(reader.pages, 1):
        contents = page.get_contents()
        expanded += len(contents.get_data()) if contents is not None else 0
        if expanded > MAX_EXPANDED_BYTES:
            raise MarkdownExportError('El PDF supera el límite de contenido descomprimido.', 'document_too_large', 413)
        text = (page.extract_text() or '').strip()
        if text:
            output.append(literal(text))
        else:
            blank_pages.append(str(number))
    warnings = [PDF_WARNING]
    if blank_pages:
        warnings.append('Páginas sin texto extraíble: ' + ', '.join(blank_pages) + '. Pueden contener imágenes o escaneos.')
    return output.render(), warnings


def word_text(node, *, escape=True):
    values = []
    for element in node.iter():
        if element.tag == WORD_NS + 'p' and values:
            values.append('\n')
        elif element.tag == WORD_NS + 't':
            values.append(element.text or '')
        elif element.tag == WORD_NS + 'tab':
            values.append('\t')
        elif element.tag in (WORD_NS + 'br', WORD_NS + 'cr'):
            values.append('\n')
    text = ''.join(values).strip()
    return literal(text) if escape else text


def word_paragraph(node, numbering):
    text = word_text(node)
    if not text:
        return ''
    style = node.find('./' + WORD_NS + 'pPr/' + WORD_NS + 'pStyle')
    name = style.get(WORD_NS + 'val', '').lower() if style is not None else ''
    if name.startswith(('heading', 'titulo', 'título')) and name[-1:].isdigit():
        return '#' * min(max(int(name[-1]), 1), 6) + ' ' + text
    if name == 'title':
        return '# ' + text
    props = node.find('./' + WORD_NS + 'pPr/' + WORD_NS + 'numPr')
    if props is not None:
        number = props.find(WORD_NS + 'numId')
        level = props.find(WORD_NS + 'ilvl')
        num_id = number.get(WORD_NS + 'val') if number is not None else ''
        indent = min(int(level.get(WORD_NS + 'val', '0')), 8) if level is not None else 0
        marker = '1. ' if numbering.get((num_id, str(indent))) != 'bullet' else '- '
        return '  ' * indent + marker + text
    return text


def word_numbering(archive):
    if 'word/numbering.xml' not in archive.namelist():
        return {}
    root = ElementTree.fromstring(archive.read('word/numbering.xml'))
    formats = {}
    for definition in root.findall(WORD_NS + 'abstractNum'):
        for level in definition.findall(WORD_NS + 'lvl'):
            fmt = level.find(WORD_NS + 'numFmt')
            formats[(definition.get(WORD_NS + 'abstractNumId'), level.get(WORD_NS + 'ilvl'))] = fmt.get(WORD_NS + 'val') if fmt is not None else 'decimal'
    result = {}
    for number in root.findall(WORD_NS + 'num'):
        abstract = number.find(WORD_NS + 'abstractNumId')
        if abstract is not None:
            for level in range(9):
                result[(number.get(WORD_NS + 'numId'), str(level))] = formats.get((abstract.get(WORD_NS + 'val'), str(level)), 'decimal')
    return result


def docx_markdown(data):
    output = MarkdownBuffer()
    with office_archive(data) as archive:
        numbering = word_numbering(archive)
        names = ['word/document.xml'] + sorted(
            name for name in archive.namelist()
            if name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes')) and name.endswith('.xml')
        )
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            body = root.find(WORD_NS + 'body') if name == 'word/document.xml' else root
            if body is None:
                continue
            for node in body:
                if node.tag == WORD_NS + 'p':
                    output.append(word_paragraph(node, numbering))
                elif node.tag == WORD_NS + 'tbl':
                    rows = [[word_text(cell) for cell in row.findall(WORD_NS + 'tc')] for row in node.findall(WORD_NS + 'tr')]
                    if rows:
                        width = max(len(row) for row in rows)
                        rows = [row + [''] * (width - len(row)) for row in rows]
                        output.append(table(rows[0], rows[1:]))
                elif node.tag in (WORD_NS + 'footnote', WORD_NS + 'endnote'):
                    output.append(word_text(node))
    return output.render(), [OFFICE_WARNING]


def xlsx_markdown(data):
    from openpyxl import load_workbook

    with office_archive(data):
        pass
    output = MarkdownBuffer()
    count = 0
    characters = 0
    warnings = [OFFICE_WARNING, 'Las fórmulas se muestran como texto; no se calculan ni se consultan vínculos externos.']
    book = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    try:
        for sheet in book.worksheets:
            rows = []
            for row in sheet.iter_rows():
                count += len(row)
                if count > MAX_SHEET_CELLS:
                    raise MarkdownExportError('El libro supera el límite de 20.000 celdas.', 'document_too_large', 413)
                values = [literal(cell.value) for cell in row]
                characters += sum(len(value) for value in values)
                if characters > 1_000_000:
                    raise MarkdownExportError('El libro supera el límite de un millón de caracteres.', 'document_too_large', 413)
                if any(values):
                    rows.append(values)
            if rows:
                output.append('## ' + literal(sheet.title))
                width = max(len(row) for row in rows)
                rows = [row + [''] * (width - len(row)) for row in rows]
                output.append(table(rows[0], rows[1:]))
    finally:
        book.close()
    return output.render(), warnings


class SourceFragmentBuffer:
    """Keep source excerpts within explicit limits without hiding omissions."""

    def __init__(self, warnings=()):
        self.fragments = []
        self.warnings = list(warnings)
        self.characters = 0
        self.partial = False
        self.limits = {
            'max_characters': MAX_SOURCE_CHARACTERS,
            'max_fragments': MAX_SOURCE_FRAGMENTS,
            'max_section_title_characters': MAX_SOURCE_SECTION_CHARACTERS,
            'max_pdf_pages': MAX_PDF_PAGES,
            'max_file_bytes': MAX_FILE_BYTES,
            'max_expanded_bytes': MAX_EXPANDED_BYTES,
        }

    def omit(self, message):
        self.partial = True
        if message not in self.warnings:
            self.warnings.append(message)

    def append(self, locator, text):
        text = text.strip()
        if not text:
            return True
        if len(self.fragments) >= MAX_SOURCE_FRAGMENTS:
            self.omit('Se alcanzó el límite de 1.000 fragmentos; quedó contenido sin incluir.')
            return False
        remaining = MAX_SOURCE_CHARACTERS - self.characters
        clipped = len(text) > remaining
        excerpt = text[:remaining]
        if excerpt:
            self.fragments.append({'locator': locator, 'text': excerpt})
            self.characters += len(excerpt)
        if clipped:
            self.omit('Se alcanzó el límite de 60.000 caracteres; quedó contenido sin incluir.')
        return not clipped

    def payload(self):
        status = 'unreadable'
        if self.fragments:
            status = 'partial' if self.partial else 'included'
        return {
            'fragments': self.fragments,
            'status': status,
            'warnings': self.warnings,
            'limits': {
                **self.limits, 'characters_returned': self.characters,
                'fragments_returned': len(self.fragments),
            },
        }


def pdf_source_fragments(data):
    """Read actual PDF pages while preserving incomplete extraction evidence."""
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    if reader.is_encrypted:
        raise MarkdownExportError('El PDF está protegido. Usa una copia sin contraseña.', 'document_protected')
    output = SourceFragmentBuffer([PDF_WARNING])
    total = len(reader.pages)
    output.limits.update({'pdf_pages_total': total, 'pdf_pages_examined': 0})
    if total > MAX_PDF_PAGES:
        output.omit(f'El PDF tiene {total} páginas; sólo se revisan las primeras 100.')
    expanded = 0
    for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], 1):
        output.limits['pdf_pages_examined'] = number
        try:
            contents = page.get_contents()
            expanded += len(contents.get_data()) if contents is not None else 0
            if expanded > MAX_EXPANDED_BYTES:
                output.omit(f'La página {number} supera el límite de contenido descomprimido; no se terminó de leer el PDF.')
                break
            text = (page.extract_text() or '').strip()
        except MemoryError:
            raise
        except Exception:
            output.omit(f'No se pudo extraer texto de la página {number}.')
            continue
        if not text:
            output.omit(f'Página {number} sin texto extraíble; puede contener imágenes o un escaneo. No se aplicó OCR.')
        elif not output.append(f'Página {number}', text):
            break
    if not output.fragments:
        output.warnings.append('El PDF no contiene texto disponible para preparar una guía; requiere revisión del original.')
    return output.payload()


def source_section_title(text, output):
    """Make any shortened section locator explicit rather than silent."""
    if len(text) > MAX_SOURCE_SECTION_CHARACTERS:
        output.omit('Un título de sección supera 200 caracteres; su referencia se abrevia y requiere revisión.')
        return text[:MAX_SOURCE_SECTION_CHARACTERS] + '…'
    return text


def _office_part_label(name):
    if name == 'word/document.xml':
        return 'Documento'
    if name.startswith('word/header'):
        return 'Encabezado ' + Path(name).stem.removeprefix('header')
    if name.startswith('word/footer'):
        return 'Pie de página ' + Path(name).stem.removeprefix('footer')
    return 'Notas al pie' if name.startswith('word/footnotes') else 'Notas finales'


def docx_source_fragments(data):
    """Locate DOCX text by XML paragraph order, never by invented pages."""
    output = SourceFragmentBuffer([OFFICE_WARNING])
    with office_archive(data) as archive:
        names = ['word/document.xml'] + sorted(
            name for name in archive.namelist()
            if name.startswith(('word/header', 'word/footer', 'word/footnotes', 'word/endnotes')) and name.endswith('.xml')
        )
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            body = root.find(WORD_NS + 'body') if name == 'word/document.xml' else root
            if body is None:
                continue
            section = ''
            for number, paragraph in enumerate(body.iter(WORD_NS + 'p'), 1):
                text = word_text(paragraph, escape=False)
                style = paragraph.find('./' + WORD_NS + 'pPr/' + WORD_NS + 'pStyle')
                style_name = style.get(WORD_NS + 'val', '').lower() if style is not None else ''
                if text and style_name.startswith(('heading', 'titulo', 'título', 'title')):
                    section = source_section_title(text, output)
                locator = _office_part_label(name)
                if section:
                    locator += f' · sección «{section}»'
                locator += f' · párrafo {number}'
                if not output.append(locator, text):
                    return output.payload()
    if not output.fragments:
        output.warnings.append('El DOCX no contiene texto extraíble; las imágenes y los escaneos requieren revisión del original.')
    return output.payload()


def extract_attachment_source(data, suffix):
    """Use the same isolated parser for page/paragraph source fragments."""
    if suffix not in ('.pdf', '.docx'):
        raise MarkdownExportError('El formato no admite extracción de fuentes verificables.', 'unsupported_format', 415)
    if len(data) > MAX_FILE_BYTES:
        raise MarkdownExportError('El archivo supera el límite de 15 MB.', 'document_too_large', 413)
    return _extract_attachment_bytes(data, suffix, 'fragments')


def _extract_attachment_bytes(data, suffix, mode='markdown'):
    try:
        result = subprocess.run(
            [sys.executable, '-m', 'content.services.attachment_markdown_worker', suffix, mode],
            input=data, capture_output=True, timeout=12,
            cwd=Path(__file__).resolve().parents[2], check=False,
        )
        if result.returncode:
            raise MarkdownExportError('La extracción superó el límite de procesamiento. Descarga el original o usa un archivo más pequeño.', 'extraction_limit', 413)
        payload = json.loads(result.stdout)
        if 'error' in payload:
            raise MarkdownExportError(payload['error'], payload['code'], payload['status'])
        return payload
    except MarkdownExportError:
        raise
    except subprocess.TimeoutExpired as exc:
        raise MarkdownExportError('La extracción tardó demasiado. Usa un archivo más pequeño.', 'extraction_limit', 413) from exc
    except Exception as exc:
        raise MarkdownExportError('No se pudo leer el archivo. Puede estar dañado o protegido.', 'invalid_document') from exc


def extract_attachment_markdown(document):
    suffix = Path(document.file.name or '').suffix.lower()
    if suffix not in ('.pdf', '.docx', '.xlsx'):
        message = 'Este formato no admite copia de texto. Las imágenes requieren reconocimiento de texto.'
        if suffix in ('.doc', '.xls'):
            message = 'Convierte el archivo a DOCX o XLSX para visualizar y copiar su contenido.'
        raise MarkdownExportError(message, 'unsupported_format', 415)
    try:
        # Parsers can expand compressed data before their page/cell limits can
        # run. Bound their memory and CPU in a short-lived process, not the web
        # worker. No shell, remote URLs, temporary files or DB access.
        payload = _extract_attachment_bytes(read_attachment(document), suffix)
        return export_payload(document.title, payload['markdown'], payload['warnings'])
    except MarkdownExportError:
        raise
    except subprocess.TimeoutExpired as exc:
        raise MarkdownExportError('La extracción tardó demasiado. Usa un archivo más pequeño.', 'extraction_limit', 413) from exc
    except Exception as exc:
        raise MarkdownExportError('No se pudo leer el archivo. Puede estar dañado o protegido.', 'invalid_document') from exc
